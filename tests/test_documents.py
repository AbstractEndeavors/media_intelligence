"""Native fixture and contract checks; OCR responses are mocked, never inferred."""

import io, json, sys, tempfile, unittest, zipfile
from pathlib import Path
from unittest.mock import patch
from media_intelligence.documents import (
    extract_document,
    read_page,
    search_document,
    Limits,
    OCRResult,
    run_job,
)
from media_intelligence.document_engine.ocr import HugpyOCR


class Documents(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.addCleanup(self.tmp.cleanup)
        self.root = Path(self.tmp.name)
        self.cache = self.root / "cache"

    def extract(self, p, **kw):
        return extract_document(p, store_root=self.cache, ocr="off", **kw)

    def text(self, name="a.txt", content="first"):
        p = self.root / name
        p.parent.mkdir(parents=True, exist_ok=True)
        p.write_text(content)
        return p

    def pdf(self, name="page.pdf", scan=False, mixed=False):
        # PDF is produced using Word-independent, native PDF objects.
        from PyPDF2 import PdfWriter
        from PyPDF2.generic import (
            DictionaryObject,
            NameObject,
            DecodedStreamObject,
            ArrayObject,
            NumberObject,
        )

        writer = PdfWriter()
        for is_scan in [False, True] if mixed else [scan]:
            writer.add_blank_page(width=400, height=400)
            page = writer.pages[-1]
            if is_scan:
                image = DecodedStreamObject()
                image.set_data(b"\xff\xff\xff" * 16)
                image.update(
                    {
                        NameObject("/Type"): NameObject("/XObject"),
                        NameObject("/Subtype"): NameObject("/Image"),
                        NameObject("/Width"): NumberObject(4),
                        NameObject("/Height"): NumberObject(4),
                        NameObject("/ColorSpace"): NameObject("/DeviceRGB"),
                        NameObject("/BitsPerComponent"): NumberObject(8),
                    }
                )
                page[NameObject("/Resources")] = DictionaryObject(
                    {
                        NameObject("/XObject"): DictionaryObject(
                            {NameObject("/Im0"): writer._add_object(image)}
                        )
                    }
                )
                stream = b"q 400 0 0 400 0 0 cm /Im0 Do Q"
            else:
                font = DictionaryObject(
                    {
                        NameObject("/Type"): NameObject("/Font"),
                        NameObject("/Subtype"): NameObject("/Type1"),
                        NameObject("/BaseFont"): NameObject("/Helvetica"),
                    }
                )
                page[NameObject("/Resources")] = DictionaryObject(
                    {
                        NameObject("/Font"): DictionaryObject(
                            {NameObject("/F1"): writer._add_object(font)}
                        )
                    }
                )
                stream = b"BT /F1 12 Tf 30 350 Td (Native invoice 8731) Tj ET"
            contents = DecodedStreamObject()
            contents.set_data(stream)
            page[NameObject("/Contents")] = writer._add_object(contents)
        p = self.root / name
        with p.open("wb") as f:
            writer.write(f)
        return p

    def test_identity_content_and_source_and_options(self):
        a = self.text("one/same.txt")
        b = self.text("two/same.txt", "second")
        x = self.extract(a)
        y = self.extract(b)
        self.assertNotEqual(x["artifact_id"], y["artifact_id"])
        self.assertEqual(y["text"], "second")
        self.assertTrue(self.extract(a)["cache_hit"])
        a.write_text("changed")
        self.assertNotEqual(x["artifact_id"], self.extract(a)["artifact_id"])
        with self.assertRaises(PermissionError):
            read_page(a, x["artifact_id"], 1, store_root=self.cache)
        self.assertNotEqual(
            self.extract(a, limits=Limits(max_pages=1))["artifact_id"],
            self.extract(a)["artifact_id"],
        )

    def test_artifact_bound_to_source_and_one_based_page(self):
        a = self.text()
        b = self.text("b.txt")
        x = self.extract(a)
        with self.assertRaises(PermissionError):
            read_page(b, x["artifact_id"], 1, store_root=self.cache)
        with self.assertRaises(ValueError):
            read_page(a, x["artifact_id"], 0, store_root=self.cache)
        self.assertEqual(
            read_page(a, x["artifact_id"], 1, store_root=self.cache)["page"]["index"], 0
        )
        self.assertEqual(
            search_document(a, x["artifact_id"], "FIRST", store_root=self.cache)[
                "hits"
            ][0]["page"],
            1,
        )
        with self.assertRaises(FileNotFoundError):
            read_page(a, "../cache", 1, store_root=self.cache)

    def test_text_encoding_and_limits(self):
        p = self.root / "unicode.txt"
        p.write_bytes("Amount €8731".encode("utf-16"))
        self.assertEqual(self.extract(p)["text"], "Amount €8731")
        p.write_text("abcdefgh")
        r = self.extract(p, limits=Limits(max_chars=3))
        self.assertFalse(r["complete"])
        self.assertTrue(r["truncated"])
        self.assertEqual(r["text"], "abc")
        self.assertEqual(r["pages"][0]["blocks"], [])
        with self.assertRaises(ValueError):
            self.extract(p, limits=Limits(max_pages=1.2))
        p.write_bytes(b"\x00Binary")
        self.assertFalse(self.extract(p)["ok"])

    def test_docx_table_header_and_merged_cells(self):
        import docx

        d = docx.Document()
        d.add_paragraph("Before")
        t = d.add_table(rows=2, cols=2)
        t.cell(0, 0).merge(t.cell(0, 1)).text = "Merged"
        t.cell(1, 0).text = "8731"
        d.add_paragraph("After")
        d.sections[0].header.paragraphs[0].text = "Header"
        p = self.root / "doc.docx"
        d.save(p)
        r = self.extract(p)
        self.assertTrue(r["complete"])
        self.assertTrue(
            all(t in r["text"] for t in ("Before", "Merged", "8731", "After", "Header"))
        )
        table = next(b for b in r["pages"][0]["blocks"] if b["type"] == "table")
        self.assertEqual(table["cells"][0]["column_span"], 2)
        self.assertEqual(r["meta"]["pagination"], "logical-document")

    def test_xlsx_formula_evidence(self):
        import openpyxl

        w = openpyxl.Workbook()
        w.active.append([2, 3, "=A1+B1"])
        p = self.root / "book.xlsx"
        w.save(p)
        r = self.extract(p)
        self.assertIn("=A1+B1", r["text"])
        self.assertTrue(r["needs_review"])
        self.assertTrue(r["complete"])
        c = r["pages"][0]["blocks"][0]["cells"][-1]
        self.assertTrue(c["cached_value_missing"])
        self.assertEqual(c["address"], "C1")

    def test_pdf_native_and_scan_coverage(self):
        p = self.pdf(mixed=True)
        r = self.extract(p)
        self.assertTrue(r["ok"])
        self.assertFalse(r["complete"])
        self.assertEqual(r["pages_total"], 2)
        self.assertEqual(r["pages"][1]["status"], "needs_ocr")
        self.assertEqual(r["pages"][0]["blocks"][0]["source"]["page"], 1)

    def test_partial_retry_reuses_good_page_and_retries_bad(self):
        p = self.pdf(mixed=True)

        class Fake:
            identity = "fake"
            calls = 0

            def recognize(self, *a, **kw):
                self.calls += 1
                if self.calls == 1:
                    raise RuntimeError("backend unavailable")
                return OCRResult("Recovered scan 987.65", "fake")

        fake = Fake()
        with patch(
            "media_intelligence.document_engine.pdf.render_page", return_value=b"PNG"
        ):
            a = extract_document(p, store_root=self.cache, backend=fake)
            self.assertFalse(a["complete"])
            self.assertEqual(a["pages"][1]["status"], "failed")
            b = extract_document(p, store_root=self.cache, backend=fake)
            self.assertTrue(b["complete"])
            self.assertTrue(b["needs_review"])
            self.assertEqual(fake.calls, 2)
            c = extract_document(p, store_root=self.cache, backend=fake)
            self.assertTrue(c["cache_hit"])
            self.assertEqual(fake.calls, 2)

    def test_empty_and_truncated_ocr_never_complete(self):
        p = self.pdf(scan=True)

        class Fake:
            identity = "fake"
            result = OCRResult("", "fake")

            def recognize(self, *a, **kw):
                return self.result

        fake = Fake()
        with patch(
            "media_intelligence.document_engine.pdf.render_page", return_value=b"PNG"
        ):
            a = extract_document(p, store_root=self.cache, backend=fake)
            self.assertFalse(a["ok"])
            self.assertFalse(a["complete"])
            fake.result = OCRResult("short", "fake", truncated=True)
            b = extract_document(p, store_root=self.cache, backend=fake)
            self.assertFalse(b["complete"])
            self.assertEqual(b["pages"][0]["status"], "partial")

    def test_page_and_ocr_budgets(self):
        p = self.pdf(mixed=True)
        r = self.extract(p, limits=Limits(max_pages=1))
        self.assertFalse(r["complete"])
        self.assertEqual(r["pages_total"], 2)
        r = extract_document(p, store_root=self.cache, limits=Limits(max_ocr_pages=0))
        self.assertFalse(r["complete"])
        self.assertEqual(r["pages"][1]["status"], "needs_ocr")

    def test_zip_limits(self):
        p = self.root / "bad.docx"
        with zipfile.ZipFile(p, "w") as z:
            z.writestr("big.xml", "x" * 5000)
        r = self.extract(p, limits=Limits(max_zip_bytes=10))
        self.assertFalse(r["ok"])
        self.assertIn("decompressed", r["error"])

    def test_native_job_isolation(self):
        p = self.text()
        r = run_job(p, store_root=self.cache, ocr="off")
        self.assertTrue(r["complete"])
        self.assertEqual(r["text"], "first")

    def test_hugpy_transport_contract_no_inference(self):
        class Response(io.BytesIO):
            def __enter__(self):
                return self

            def __exit__(self, *a):
                self.close()

        def reply(req, timeout):
            body = json.loads(req.data)
            self.assertEqual(body["model"], "test-vision")
            self.assertFalse(body["stream"])
            self.assertIn(
                "data:image/png;base64,",
                body["messages"][0]["content"][1]["image_url"]["url"],
            )
            return Response(
                json.dumps(
                    {
                        "choices": [
                            {
                                "message": {"content": "8731 [unreadable]"},
                                "finish_reason": "length",
                            }
                        ]
                    }
                ).encode()
            )

        with patch("urllib.request.urlopen", side_effect=reply):
            r = HugpyOCR(model="test-vision", base_url="http://mock").recognize(
                b"PNG", "image/png", timeout=10
            )
        self.assertTrue(r.truncated)
        self.assertEqual(len(r.warnings), 2)


if __name__ == "__main__":
    unittest.main()


class Presentation(unittest.TestCase):
    def test_slide_relationship_and_shape_order(self):
        with tempfile.TemporaryDirectory() as folder:
            path = Path(folder) / "slides.pptx"
            p = "http://schemas.openxmlformats.org/presentationml/2006/main"
            a = "http://schemas.openxmlformats.org/drawingml/2006/main"
            r = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
            with zipfile.ZipFile(path, "w") as z:
                z.writestr(
                    "ppt/presentation.xml",
                    f'<p:presentation xmlns:p="{p}" xmlns:r="{r}"><p:sldIdLst><p:sldId r:id="r2"/><p:sldId r:id="r1"/></p:sldIdLst></p:presentation>',
                )
                z.writestr(
                    "ppt/_rels/presentation.xml.rels",
                    '<Relationships><Relationship Id="r1" Target="slides/slide1.xml"/><Relationship Id="r2" Target="slides/slide2.xml"/></Relationships>',
                )
                z.writestr(
                    "ppt/slides/slide1.xml",
                    f'<p:sld xmlns:p="{p}" xmlns:a="{a}"><p:sp><a:p><a:t>Last slide</a:t></a:p></p:sp></p:sld>',
                )
                z.writestr(
                    "ppt/slides/slide2.xml",
                    f'<p:sld xmlns:p="{p}" xmlns:a="{a}"><p:sp><a:p><a:t>Before</a:t></a:p></p:sp><p:graphicFrame><a:tbl><a:tr><a:tc><a:p><a:t>Table 8731</a:t></a:p></a:tc></a:tr></a:tbl></p:graphicFrame><p:sp><a:p><a:t>After</a:t></a:p></p:sp></p:sld>',
                )
            result = extract_document(
                path, ocr="off", store_root=Path(folder) / "cache"
            )
            self.assertTrue(result["complete"])
            self.assertEqual(
                [b["text"] for b in result["pages"][0]["blocks"]],
                ["Before", "Table 8731", "After"],
            )
            self.assertEqual(result["pages"][1]["text"], "Last slide")


class OCRQualityChecks(unittest.TestCase):
    def test_repetition_is_review_signal(self):
        from media_intelligence.document_engine.ocr import repetition_suspect

        passage = "First second third fourth fifth sixth seventh eighth ninth tenth eleventh twelfth."
        self.assertTrue(repetition_suspect(" ".join([passage] * 3)))
        self.assertFalse(
            repetition_suspect("A brief ordinary invoice with total 987.65.")
        )

    def test_repeated_ocr_stays_available_with_warning(self):
        class Response(io.BytesIO):
            def __enter__(self):
                return self

            def __exit__(self, *args):
                self.close()

        text = " ".join(
            [
                "First second third fourth fifth sixth seventh eighth ninth tenth eleventh twelfth."
            ]
            * 3
        )
        response = Response(
            json.dumps(
                {"choices": [{"message": {"content": text}, "finish_reason": "stop"}]}
            ).encode()
        )
        with patch("urllib.request.urlopen", return_value=response):
            result = HugpyOCR(model="fake").recognize(b"PNG", "image/png", timeout=10)
        self.assertEqual(result.text, text)
        self.assertTrue(any("repeated" in w for w in result.warnings))
