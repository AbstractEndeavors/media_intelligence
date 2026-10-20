"""Native document readers, with structure and source-native references."""

from __future__ import annotations
import csv, io, zipfile, xml.etree.ElementTree as ET
from pathlib import Path
from .model import Page, Limits

TEXT_EXT = {
    ".txt",
    ".md",
    ".markdown",
    ".rst",
    ".csv",
    ".tsv",
    ".json",
    ".log",
    ".text",
    ".xml",
    ".yaml",
    ".yml",
}


def check_zip(path: Path, limits: Limits):
    with zipfile.ZipFile(path) as z:
        info = z.infolist()
        if len(info) > 10000 or sum(i.file_size for i in info) > limits.max_zip_bytes:
            raise ValueError(
                "Document archive exceeds the decompressed-size or entry limit"
            )
        if any(i.flag_bits & 1 for i in info):
            raise ValueError("Encrypted Office archives are not supported")


def text_page(path: Path, limits: Limits) -> Page:
    raw = path.read_bytes()
    warnings = []
    if raw.startswith((b"\xff\xfe\x00\x00", b"\x00\x00\xfe\xff")):
        encoding = "utf-32"
    elif raw.startswith((b"\xff\xfe", b"\xfe\xff")):
        encoding = "utf-16"
    elif raw.startswith(b"\xef\xbb\xbf"):
        encoding = "utf-8-sig"
    else:
        encoding = "utf-8"
    try:
        text = raw.decode(encoding)
    except UnicodeDecodeError:
        if b"\x00" in raw:
            raise ValueError("Binary content cannot be read as a text document")
        encoding = "cp1252"
        try:
            text = raw.decode(encoding)
        except UnicodeDecodeError:
            raise ValueError("Text encoding is not recognized; save as UTF-8")
        warnings.append(
            "UTF-8 decoding failed; text was decoded as Windows-1252 and needs review"
        )
    if "\x00" in text:
        raise ValueError(
            "Text contains NUL characters; its encoding is unsupported or the file is binary"
        )
    page = Page(
        0,
        text.strip(),
        method="native-text:" + encoding,
        warnings=warnings,
        needs_review=bool(warnings),
    )
    page.blocks = [
        {
            "type": "text",
            "text": page.text,
            "source": {"line_start": 1, "line_end": len(text.splitlines())},
        }
    ]
    if path.suffix.lower() in (".csv", ".tsv"):
        rows = []
        count = 0
        for row in csv.reader(
            io.StringIO(text), delimiter="\t" if path.suffix.lower() == ".tsv" else ","
        ):
            count += len(row)
            if count > limits.max_cells:
                page.status = "partial"
                page.warnings.append("Delimited table exceeded the cell limit")
                break
            rows.append(row)
        page.blocks = [
            {
                "type": "table",
                "rows": rows,
                "text": page.text,
                "source": {"line_start": 1},
            }
        ]
    return page


def _paragraph(p, source):
    text = p.text.strip()
    if not text:
        return None
    style = p.style.name if p.style else ""
    level = None
    if style.startswith("Heading "):
        try:
            level = int(style.split()[-1])
        except ValueError:
            pass
    return {
        "type": "heading" if level else "paragraph",
        "text": text,
        "level": level,
        "source": source,
    }


def _word_table(table, source):
    cells = []
    seen = {}
    for r, row in enumerate(table.rows):
        for c, cell in enumerate(row.cells):
            key = cell._tc
            if key in seen:
                item = seen[key]
                item["row_span"] = max(item["row_span"], r - item["row"] + 1)
                item["column_span"] = max(item["column_span"], c - item["column"] + 1)
                continue
            item = {
                "row": r,
                "column": c,
                "row_span": 1,
                "column_span": 1,
                "text": cell.text,
                "source": {**source, "row": r, "column": c},
            }
            if cell.tables:
                item["tables"] = [
                    _word_table(t, {**source, "row": r, "column": c, "nested_table": i})
                    for i, t in enumerate(cell.tables)
                ]
            seen[key] = item
            cells.append(item)
    lines = [
        " | ".join(cell.text.replace("\n", " / ") for cell in row.cells)
        for row in table.rows
    ]
    return {"type": "table", "text": "\n".join(lines), "cells": cells, "source": source}


def docx_pages(path: Path, limits: Limits):
    check_zip(path, limits)
    import docx
    from docx.table import Table

    d = docx.Document(path)
    blocks = []
    warnings = []
    body_chars = 0
    if not hasattr(d, "iter_inner_content"):
        raise RuntimeError("Document reading requires python-docx >=1.2")
    for index, item in enumerate(d.iter_inner_content()):
        source = {"part": "body", "index": index}
        block = (
            _word_table(item, source)
            if isinstance(item, Table)
            else _paragraph(item, source)
        )
        if block:
            blocks.append(block)
            body_chars += len(block.get("text", ""))
        if body_chars > limits.max_chars:
            break
    seen = set()
    for section_index, section in enumerate(d.sections):
        for label in (
            "header",
            "first_page_header",
            "even_page_header",
            "footer",
            "first_page_footer",
            "even_page_footer",
        ):
            part = getattr(section, label)
            if part.is_linked_to_previous or str(part.part.partname) in seen:
                continue
            seen.add(str(part.part.partname))
            for index, item in enumerate(part.iter_inner_content()):
                source = {"part": label, "section": section_index, "index": index}
                block = (
                    _word_table(item, source)
                    if isinstance(item, Table)
                    else _paragraph(item, source)
                )
                if block:
                    blocks.append(block)
    # Footnotes/endnotes have no high-level python-docx traversal API.
    with zipfile.ZipFile(path) as z:
        for part in ("footnotes", "endnotes"):
            filename = "word/" + part + ".xml"
            if filename not in z.namelist():
                continue
            xml = ET.fromstring(z.read(filename))
            ns = {"w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main"}
            for node in xml:
                text = "".join(t.text or "" for t in node.findall(".//w:t", ns)).strip()
                if text:
                    blocks.append(
                        {
                            "type": "note",
                            "text": text,
                            "source": {
                                "part": part,
                                "id": node.get("{" + ns["w"] + "}id"),
                            },
                        }
                    )
    page = Page(
        0,
        "\n\n".join(b["text"] for b in blocks),
        blocks=blocks,
        method="native-docx",
        warnings=warnings,
    )
    # Logical document section: never claim a Word source has stable PDF-style pages.
    page.name = "Document"
    return [page], {
        "pagination": "logical-document",
        "title": d.core_properties.title or None,
    }


def xlsx_pages(path: Path, limits: Limits):
    check_zip(path, limits)
    import openpyxl

    formulas = openpyxl.load_workbook(path, read_only=True, data_only=False)
    values = openpyxl.load_workbook(path, read_only=True, data_only=True)
    pages = []
    count = 0
    try:
        total = len(formulas.worksheets)
        for index, ws in enumerate(formulas.worksheets[: limits.max_pages]):
            cells = []
            lines = []
            uncached = []
            truncated = False
            for row, cached_row in zip(ws.iter_rows(), values[ws.title].iter_rows()):
                line = []
                for cell, cached in zip(row, cached_row):
                    count += 1
                    if count > limits.max_cells:
                        truncated = True
                        break
                    if cell.value is None:
                        line.append("")
                        continue
                    formula = cell.value if cell.data_type == "f" else None
                    value = cached.value if formula else cell.value
                    if formula and value is None:
                        uncached.append(cell.coordinate)
                    display = str(value) if value is not None else str(formula or "")
                    line.append(display)
                    cells.append(
                        {
                            "address": cell.coordinate,
                            "value": value.isoformat()
                            if hasattr(value, "isoformat")
                            else value,
                            "formula": formula,
                            "cached_value_missing": bool(formula and value is None),
                        }
                    )
                if any(line):
                    lines.append("\t".join(line).rstrip("\t"))
                if truncated:
                    break
            warnings = []
            if uncached:
                warnings.append(
                    f"{len(uncached)} formula cells have no cached value; formulas are preserved without calculating them: "
                    + ", ".join(uncached[:10])
                )
            if truncated:
                warnings.append("Workbook exceeded the cell limit")
            text = "# Sheet: " + ws.title + "\n" + "\n".join(lines)
            pages.append(
                Page(
                    index,
                    text,
                    method="native-xlsx",
                    name=ws.title,
                    blocks=[
                        {
                            "type": "table",
                            "text": text,
                            "cells": cells,
                            "source": {"sheet": ws.title},
                        }
                    ],
                    warnings=warnings,
                    needs_review=bool(uncached),
                    status="partial" if truncated else "complete",
                )
            )
            if truncated:
                break
        return pages, {"pagination": "worksheet", "page_count": total}
    finally:
        formulas.close()
        values.close()


def pptx_pages(path: Path, limits: Limits):
    check_zip(path, limits)
    ns = {
        "a": "http://schemas.openxmlformats.org/drawingml/2006/main",
        "p": "http://schemas.openxmlformats.org/presentationml/2006/main",
    }
    with zipfile.ZipFile(path) as z:
        pres = ET.fromstring(z.read("ppt/presentation.xml"))
        rels = ET.fromstring(z.read("ppt/_rels/presentation.xml.rels"))
        targets = {
            r.get("Id"): r.get("Target")
            for r in rels
            if r.get("TargetMode") != "External"
        }
        paths = []
        for node in pres.findall(".//p:sldId", ns):
            target = targets.get(
                node.get(
                    "{http://schemas.openxmlformats.org/officeDocument/2006/relationships}id"
                )
            )
            if not target:
                continue
            import posixpath

            resolved = posixpath.normpath(posixpath.join("ppt", target))
            if not resolved.startswith("ppt/slides/"):
                raise ValueError("Invalid slide relationship")
            paths.append(resolved)
        pages = []
        for i, name in enumerate(paths[: limits.max_pages]):
            xml = ET.fromstring(z.read(name))
            blocks = []
            for index, node in enumerate(
                [
                    n
                    for n in xml.iter()
                    if n.tag in ("{" + ns["p"] + "}sp", "{" + ns["p"] + "}graphicFrame")
                ]
            ):
                rows = node.findall(".//a:tr", ns)
                if rows:
                    matrix = [
                        [
                            " ".join(t.text or "" for t in c.findall(".//a:t", ns))
                            for c in row.findall("a:tc", ns)
                        ]
                        for row in rows
                    ]
                    blocks.append(
                        {
                            "type": "table",
                            "rows": matrix,
                            "text": "\n".join(" | ".join(r) for r in matrix),
                            "source": {"slide": i + 1, "shape": index},
                        }
                    )
                else:
                    text = "\n".join(
                        "".join(t.text or "" for t in p.findall(".//a:t", ns))
                        for p in node.findall(".//a:p", ns)
                    ).strip()
                    if text:
                        blocks.append(
                            {
                                "type": "paragraph",
                                "text": text,
                                "source": {"slide": i + 1, "shape": index},
                            }
                        )
            pages.append(
                Page(
                    i,
                    "\n\n".join(b["text"] for b in blocks),
                    blocks=blocks,
                    method="native-pptx",
                    name="Slide " + str(i + 1),
                )
            )
        return pages, {"pagination": "slide", "page_count": len(paths)}
